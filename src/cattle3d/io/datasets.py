"""Adapters for the two public Kinect v2 herds.

Owns: how each dataset lays out its files and its measurement table, and the
join between table rows and point-cloud files. Must not know about
segmentation, features or models.

Herd A, CowDatabase2 (github.com/ruchaya/CowDatabase2): Black Angus, one folder
per capture frame, ``<frame>/raw_data/<view>/cloud.ply``. The table
``Database2022v2.xlsx`` joins to folders by its *Frame number* column, not by
the animal number. Rows whose frame is not an integer have no folder.

Herd B, CowDatabase (github.com/ruchaya/CowDatabase): Hereford,
``<View>Clouds/<file>.ply`` with ``Measurements.xlsx`` keyed by ``N``. The
published release assumes file number == N, but for files 22-87 that join is
off by one or two rows (see reports/cowdb1_join_audit.md). The audited join in
``resources/cowdb1_audited_join.csv`` is used by default; it was derived from
height measurements only, never from weight.
"""

from __future__ import annotations

import csv
from dataclasses import dataclass
from importlib import resources
from pathlib import Path
from typing import Literal

import openpyxl

from cattle3d.types import ALL_VIEWS, AnimalRecord, View

HERD_A = "cowdb2"
HERD_B = "cowdb1"

# Manual measurements present in both herds, in centimetres. These are the
# comparator for the "volume beats manual zoometry" hypothesis.
COMMON_MANUAL = ("withers_height", "hip_height", "chest_width", "heart_girth")


class DatasetLayoutError(RuntimeError):
    """The dataset folder does not look like the expected public release."""


@dataclass(frozen=True)
class SkippedRow:
    animal_id: str
    reason: str


@dataclass(frozen=True)
class LoadedHerd:
    herd: str
    animals: tuple[AnimalRecord, ...]
    skipped: tuple[SkippedRow, ...]


def _as_float(value: object) -> float | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        try:
            return float(value.strip().replace(",", "."))
        except ValueError:
            return None
    return None


def _missing_views(paths: dict[View, Path]) -> list[str]:
    return [view.value for view, path in paths.items() if not path.is_file()]


def load_cowdatabase2(root: Path) -> LoadedHerd:
    """Load herd A (CowDatabase2)."""
    root = Path(root)
    table = root / "Database2022v2.xlsx"
    if not table.is_file():
        raise DatasetLayoutError(f"{table} not found; is {root} a CowDatabase2 checkout?")

    sheet = openpyxl.load_workbook(table, read_only=True, data_only=True).worksheets[0]
    rows = list(sheet.iter_rows(values_only=True))
    header = [str(c).strip() if c is not None else "" for c in rows[0]]
    try:
        col_frame = header.index("Frame number")
        col_weight = header.index("Live weigth")  # sic, as spelled in the release
        col_withers = header.index("Withers height")
        col_hip = header.index("Hip height")
        col_chest_w = header.index("Chest width")
        col_girth = header.index("Chest heart")
    except ValueError as exc:
        raise DatasetLayoutError(f"unexpected CowDatabase2 header: {exc}") from exc

    animals: list[AnimalRecord] = []
    skipped: list[SkippedRow] = []
    for row in rows[1:]:
        weight = _as_float(row[col_weight])
        frame = _as_float(row[col_frame])
        if weight is None or frame is None:
            continue  # header continuation or empty line
        if frame != int(frame):
            skipped.append(SkippedRow(str(frame), "frame number is not an integer; no folder"))
            continue
        animal_id = str(int(frame))
        paths = {v: root / animal_id / "raw_data" / v.value / "cloud.ply" for v in ALL_VIEWS}
        missing = _missing_views(paths)
        if missing:
            skipped.append(SkippedRow(animal_id, f"missing views: {', '.join(missing)}"))
            continue
        manual = {
            "withers_height": _as_float(row[col_withers]),
            "hip_height": _as_float(row[col_hip]),
            "chest_width": _as_float(row[col_chest_w]),
            "heart_girth": _as_float(row[col_girth]),
        }
        animals.append(
            AnimalRecord(
                herd=HERD_A,
                animal_id=animal_id,
                weight_kg=weight,
                view_paths=paths,
                manual={k: v for k, v in manual.items() if v is not None},
            )
        )
    return LoadedHerd(HERD_A, tuple(animals), tuple(skipped))


_COWDB1_COLUMNS = {
    "live weithg": "weight",  # sic
    "withers height": "withers_height",
    "hip height": "hip_height",
    "chest depth": "chest_depth",
    "chest width": "chest_width",
    "ilium width": "ilium_width",
    "hip joint width": "hip_joint_width",
    "oblique body length": "oblique_body_length",
    "hip length": "hip_length",
    "heart girth": "heart_girth",
}


Cowdb1Join = Literal["audited", "published"]


def audited_cowdb1_join() -> dict[str, str]:
    """Table row ``N`` -> cloud file number, from the packaged audit result."""
    text = resources.files("cattle3d.resources").joinpath("cowdb1_audited_join.csv").read_text()
    return {row["row_n"]: row["file_id"] for row in csv.DictReader(text.splitlines())}


def load_cowdatabase(root: Path, join: Cowdb1Join = "audited") -> LoadedHerd:
    """Load herd B (CowDatabase).

    ``join="published"`` pairs table row N with file N, as the release does;
    ``join="audited"`` uses the corrected pairing and drops rows without a
    matching capture.
    """
    root = Path(root)
    table = root / "Measurements.xlsx"
    if not table.is_file():
        raise DatasetLayoutError(f"{table} not found; is {root} a CowDatabase checkout?")

    sheet = openpyxl.load_workbook(table, read_only=True, data_only=True).worksheets[0]
    rows = list(sheet.iter_rows(values_only=True))
    header = [str(c).strip().lower() if c is not None else "" for c in rows[0]]
    try:
        col_n = header.index("n")
        cols = {name: header.index(src) for src, name in _COWDB1_COLUMNS.items()}
    except ValueError as exc:
        raise DatasetLayoutError(f"unexpected CowDatabase header: {exc}") from exc

    file_for_row = audited_cowdb1_join() if join == "audited" else None
    animals: list[AnimalRecord] = []
    skipped: list[SkippedRow] = []
    for row in rows[1:]:
        n = _as_float(row[col_n])
        weight = _as_float(row[cols["weight"]])
        if n is None or weight is None:
            continue
        animal_id = str(int(n))
        file_id = animal_id if file_for_row is None else file_for_row.get(animal_id)
        if file_id is None:
            skipped.append(SkippedRow(animal_id, "no capture matches this row in the audited join"))
            continue
        paths = {v: root / f"{v.value.capitalize()}Clouds" / f"{file_id}.ply" for v in ALL_VIEWS}
        missing = _missing_views(paths)
        if missing:
            skipped.append(SkippedRow(animal_id, f"missing views: {', '.join(missing)}"))
            continue
        manual = {
            name: value
            for name, col in cols.items()
            if name != "weight" and (value := _as_float(row[col])) is not None
        }
        animals.append(AnimalRecord(HERD_B, animal_id, weight, paths, manual))
    return LoadedHerd(HERD_B, tuple(animals), tuple(skipped))


LOADERS = {HERD_A: load_cowdatabase2, HERD_B: load_cowdatabase}
